#!/usr/bin/env python3
"""Prepare Cadentra's catalog-only UCI subset from an explicitly supplied archive.

This offline preparation command requires openpyxl; it is not an application
runtime dependency. It never connects to an application/database or imports
customer, contact, price or historical transaction fields.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

SOURCE_URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
DATASET_URL = "https://archive.ics.uci.edu/dataset/352/online%2Bretail"
WORKBOOK_NAME = "Online Retail.xlsx"
STOCK_CODE = re.compile(r"[0-9]{5}[A-Z]*\Z")
NON_PRODUCT = re.compile(
    r"\b(?:postage|carriage|discount|manual|commission|fee|charges?|adjustments?|"
    r"damaged?|damages|faulty|broken|missing|lost|crushed|check|found)\b",
    re.IGNORECASE,
)


def sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def utc_timestamp(value: str) -> str:
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        raise argparse.ArgumentTypeError("Retrieval timestamp must include its UTC offset")
    return parsed.astimezone(UTC).isoformat(timespec="seconds").replace("+00:00", "Z")


def prepare(archive: Path, retrieved_at: str, count: int = 80) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("Offline preparation requires openpyxl; runtime does not.") from exc

    archive_bytes = archive.read_bytes()
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as source:
        if source.namelist() != [WORKBOOK_NAME]:
            raise ValueError("Unexpected UCI source archive contents")
        workbook_bytes = source.read(WORKBOOK_NAME)

    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    descriptions: dict[str, Counter[str]] = defaultdict(Counter)
    try:
        sheet = workbook.active
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if headers[1:3] != ("StockCode", "Description"):
            raise ValueError("Expected StockCode/Description source columns B/C")
        # Retain only columns B/C. Never collect CustomerID, Country, InvoiceNo,
        # InvoiceDate, Quantity or UnitPrice in the prepared data.
        for raw_code, raw_description in sheet.iter_rows(
            min_row=2, min_col=2, max_col=3, values_only=True
        ):
            if isinstance(raw_code, bool) or not isinstance(raw_code, (str, int)):
                continue
            code = str(raw_code).strip()
            if not STOCK_CODE.fullmatch(code) or not isinstance(raw_description, str):
                continue
            original = raw_description.strip()
            if (
                not 6 <= len(original) <= 200
                or sum(char.isalpha() for char in original) < 6
                or "?" in original
                or any(ord(char) < 32 for char in original)
                or NON_PRODUCT.search(original)
            ):
                continue
            descriptions[code][original] += 1
    finally:
        workbook.close()

    ranked = sorted(descriptions, key=lambda code: (-sum(descriptions[code].values()), code))
    if len(ranked) < count:
        raise ValueError("Source contains fewer eligible products than requested")
    products = []
    for code in ranked[:count]:
        # Modal trimmed source description; exact lexical tie-break is stable.
        original = min(descriptions[code], key=lambda value: (-descriptions[code][value], value))
        products.append(
            {
                "stock_code": code,
                "original_description": original,
                "name": " ".join(original.split()).title(),
            }
        )
    if len(products) != count or len({row["stock_code"] for row in products}) != count:
        raise ValueError("Prepared product count/identity invariant failed")
    if not all(row["name"] and row["original_description"] for row in products):
        raise ValueError("Prepared product names must not be empty")

    return {
        "schema_version": 1,
        "dataset_id": "cadentra-commerce-catalog-v1",
        "source": {
            "title": "Online Retail",
            "creator": "Daqing Chen",
            "year": 2015,
            "doi": "10.24432/C5BW33",
            "dataset_url": DATASET_URL,
            "download_url": SOURCE_URL,
            "license": "CC BY 4.0",
            "license_url": "https://creativecommons.org/licenses/by/4.0/",
            "retrieved_at_utc": retrieved_at,
            "archive_sha256": sha256(archive_bytes),
            "workbook_name": WORKBOOK_NAME,
            "workbook_sha256": sha256(workbook_bytes),
        },
        "attribution": (
            "Catalog descriptions adapted from Chen, D. (2015), Online Retail, "
            "UCI Machine Learning Repository, DOI:10.24432/C5BW33. Licensed CC BY 4.0. "
            "Descriptions were selected and normalized; prices, customers and financial "
            "scenarios are fictional Cadentra demo data. No endorsement is implied."
        ),
        "selection": {
            "count": count,
            "source_columns_retained": ["StockCode", "Description"],
            "stock_code_pattern": "[0-9]{5}[A-Z]*",
            "excluded": (
                "Blank, short, oversized, control-character/question-mark descriptions; "
                "adjustment, service, fee and damaged/missing-stock descriptions. "
                "The exact exclusion expression is frozen in the preparation script."
            ),
            "ranking": (
                "Descending count of eligible description occurrences per stock code, "
                "then ascending stock code; no quantity or sales-value ranking."
            ),
            "canonical_description": (
                "Most frequent trimmed nonblank eligible description per stock code; "
                "ties use ascending exact source-description text."
            ),
            "display_normalization": "Collapse whitespace and apply Python str.title().",
            "generated_data_notice": (
                "All INR prices, costs, variants, plans, customers, subscriptions, invoices "
                "and payments are synthetic demo scenarios. Historical sterling prices "
                "are neither imported nor converted. No source customer or order fields "
                "are retained."
            ),
        },
        "products": products,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-archive", required=True, type=Path)
    parser.add_argument("--retrieved-at", required=True, type=utc_timestamp)
    parser.add_argument("--count", type=int, default=80, choices=range(1, 1001), metavar="1..1000")
    parser.add_argument(
        "--output", type=Path, default=Path("backend/app/demo_data/online_retail_catalog.json")
    )
    args = parser.parse_args()
    manifest = prepare(args.source_archive, args.retrieved_at, args.count)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Prepared {len(manifest['products'])} catalog products; no customer/order data.")
    print(f"Archive SHA-256: {manifest['source']['archive_sha256']}")
    print(f"Workbook SHA-256: {manifest['source']['workbook_sha256']}")


if __name__ == "__main__":
    main()
